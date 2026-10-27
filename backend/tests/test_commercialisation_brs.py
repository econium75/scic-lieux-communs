"""Commercialisation BRS : le moteur reproduit le classeur « Simulation BRS » de la SCIC."""
import pytest
from fastapi.testclient import TestClient
from io import BytesIO
from openpyxl import load_workbook

from app.engines import commercialisation as brs
from app.main import app

# Ménage d'exemple du classeur (onglet « Simulation ménage »)
MENAGE = {"reference": "EX", "typologie": "T3", "adultes": 2, "enfants": 1, "rfr_n2": 42_000,
          "revenu_disponible_annuel": 38_000, "revenus_nets_mensuels": 3_200, "credits_mensuels": 0, "apport": 15_000}


def test_menage_du_classeur():
    r = brs.simuler_menage(brs.hypotheses(None), MENAGE)
    assert r["personnes"] == 3 and r["uc"] == pytest.approx(1.8)
    assert r["revenu_par_uc"] == pytest.approx(21_111.11, abs=0.01) and r["tranche"] == "7-9"
    assert (r["surface"], r["prix"], r["subvention"], r["prix_net"]) == (65, 271_600, 8_200, 263_400)
    assert r["frais"] == 8_148 and r["emprunt"] == 256_548
    assert r["mensualite"] == pytest.approx(1_284.34, abs=0.01)
    assert r["assurance"] == pytest.approx(64.14, abs=0.01)
    assert (r["redevance"], r["charges"], r["taxe_fonciere"]) == (97.5, 162.5, 50)
    assert r["cout_mensuel"] == pytest.approx(1_658.48, abs=0.01)
    assert r["taux_effort"] == pytest.approx(0.4519, abs=0.0001)
    assert r["reste_a_vivre"] == pytest.approx(1_541.52, abs=0.01)
    assert r["reste_a_vivre_uc"] == pytest.approx(856.40, abs=0.01)
    # éligible (42 000 € ≤ plafond 3 personnes) mais effort de 45 % > 35 % : à retravailler
    assert r["eligible"] is True and r["verdict"] == "a_retravailler"


def test_plafond_non_saisi_donne_a_completer_et_non_pas_non_eligible():
    h = brs.hypotheses({"plafonds": {"montants": [None] * 5}})
    r = brs.simuler_menage(h, MENAGE)
    assert r["eligible"] is None and r["verdict"] == "a_completer"


def test_verdicts_financable_et_non_eligible():
    h = brs.hypotheses(None)
    ok = brs.simuler_menage(h, {**MENAGE, "typologie": "T2", "adultes": 1, "enfants": 0, "rfr_n2": 30_000,
                                "revenu_disponible_annuel": 26_000, "revenus_nets_mensuels": 2_900, "apport": 30_000})
    assert ok["tranche"] == "7-9" and ok["verdict"] == "financable", ok["motifs"]
    trop = brs.simuler_menage(h, {**MENAGE, "rfr_n2": 90_000})
    assert trop["eligible"] is False and trop["verdict"] == "non_eligible"
    hors = brs.simuler_menage(h, {**MENAGE, "revenu_disponible_annuel": 70_000})
    assert hors["tranche"] == "hors" and hors["subvention"] == 0


def test_programme_et_subventions():
    p = brs.programme(brs.hypotheses(None))
    assert [l["subvention"] for l in p["lignes"]] == [88_700, 141_200, 133_000]
    assert p["total"]["nb"] == 15 and p["total"]["subvention"] == 362_900
    assert p["total"]["prix_brut"] == 3_834_600 and p["total"]["prix_net"] == 3_471_700
    assert p["quotites_reelles"]["3-4"] == pytest.approx(0.4) and p["quotites_reelles"]["7-9"] == pytest.approx(4 / 15)
    assert p["subvention_moyenne"] == pytest.approx(24_193.33, abs=0.01) and p["coherent"]
    h = brs.hypotheses(None)
    h["typologies"][0]["repartition"]["3-4"] = 3
    assert brs.programme(h)["lignes"][0]["controle"] == "Écart +1"


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def _h(client, user, pwd):
    return {"Authorization": f"Bearer {client.post('/api/auth/login', data={'username': user, 'password': pwd}).json()['access_token']}"}


def test_api_hypotheses_menages_export(client):
    saisie, lecture = _h(client, "saisie", "saisie123"), _h(client, "lecture", "lecture123")
    sid = client.get("/api/operations/1/scenarios").json()[0]["id"]
    assert client.get(f"/api/brs/scenario/{sid}").status_code == 401
    e = client.get(f"/api/brs/scenario/{sid}", headers=lecture).json()
    assert e["programme"]["total"]["nb"] == 15
    assert client.post(f"/api/brs/scenario/{sid}/menages", headers=lecture, json=MENAGE).status_code == 403
    e = client.post(f"/api/brs/scenario/{sid}/menages", headers=saisie, json=MENAGE).json()
    m = next(x for x in e["menages"] if x["reference"] == "EX")
    assert m["resultat"]["verdict"] == "a_retravailler"
    # un apport plus élevé et une durée plus longue rendent le projet finançable
    h = {**e["hypotheses"], "duree_ans": 30}
    assert client.put(f"/api/brs/scenario/{sid}/hypotheses", headers=saisie, json={**h, "taux_pret": 3.5}).status_code == 422
    client.put(f"/api/brs/scenario/{sid}/hypotheses", headers=saisie, json=h)
    e = client.put(f"/api/brs/menages/{m['id']}", headers=saisie, json={"apport": 120_000}).json()
    assert next(x for x in e["menages"] if x["id"] == m["id"])["resultat"]["verdict"] == "financable"
    sim = client.post(f"/api/brs/scenario/{sid}/simuler", headers=lecture, json=MENAGE).json()
    assert sim["typologie"] == "T3" and sim["emprunt"] == 256_548
    x = client.get(f"/api/brs/scenario/{sid}/export.xlsx", headers=lecture)
    assert x.status_code == 200
    wb = load_workbook(BytesIO(x.content))
    assert wb.sheetnames == ["Hypothèses", "Ménages candidats", "Programme & subventions"]
    assert client.delete(f"/api/brs/menages/{m['id']}", headers=saisie).status_code == 200
