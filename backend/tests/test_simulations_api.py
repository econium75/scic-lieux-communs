"""API du simulateur : lecture, droits, sauvegarde, limites, exports."""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def _token(client, user, pwd):
    r = client.post("/api/auth/login", data={"username": user, "password": pwd})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def test_lecture_orly_econhome(client):
    ops = client.get("/api/operations").json()
    orly = next(o for o in ops if o["code"] == "ORLY-ECO-2027")
    r = client.get(f"/api/simulations/operation/{orly['id']}")
    assert r.status_code == 200
    body = r.json()
    assert body["data"]["hypotheses"]["nom_operation"] == "ECON'HOME"
    assert round(body["result"]["prix_revient"]["ttc"]["PLAI"], 2) == 636233.84


def test_lecteur_ne_peut_pas_modifier(client):
    h = _token(client, "lecture", "lecture123")
    sim = client.get("/api/simulations/operation/1").json()
    r = client.put("/api/simulations/operation/1", json={"data": sim["data"]}, headers=h)
    assert r.status_code == 403


def test_editeur_enregistre_et_operation_synchronisee(client):
    h = _token(client, "saisie", "saisie123")
    sim = client.get("/api/simulations/operation/1").json()
    data = sim["data"]
    data["hypotheses"]["vacance"] = 0.03
    r = client.put("/api/simulations/operation/1", json={"data": data}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["data"]["hypotheses"]["vacance"] == 0.03
    op = client.get("/api/operations/1").json()
    assert op["nb_logements_total"] == 30


def test_limite_cinq_financements_complementaires(client):
    h = _token(client, "saisie", "saisie123")
    data = client.get("/api/simulations/operation/1").json()["data"]
    data["financement"]["complementaires"] = [
        {"libelle": f"F{i}", "nature": "subvention", "montants": {"PLAI": 1}} for i in range(6)
    ]
    r = client.put("/api/simulations/operation/1", json={"data": data}, headers=h)
    assert r.status_code == 422


def test_calcul_apercu_sans_enregistrer(client):
    data = client.get("/api/simulations/operation/1").json()["data"]
    data["hypotheses"]["travaux_m2"] = 9999
    r = client.post("/api/simulations/operation/1/calcul", json={"data": data})
    assert r.status_code == 200
    assert "synthese" in r.json()


def test_creation_operation_initialise_simulateur(client):
    h = _token(client, "saisie", "saisie123")
    r = client.post("/api/operations", json={"code": "TEST-SIM", "nom": "Test simulateur", "commune": "Lyon",
                                            "zone_123": "Z2", "zone_abc": "B1", "annee_livraison": 2030}, headers=h)
    assert r.status_code == 200, r.text
    op_id = r.json()["id"]
    sim = client.get(f"/api/simulations/operation/{op_id}").json()
    assert sim["data"]["hypotheses"]["zone_abc"] == "B1"
    assert sim["data"]["hypotheses"]["annee_livraison"] == 2030
    r = client.post(f"/api/simulations/operation/{op_id}/appliquer-bareme", json={"data": sim["data"]})
    assert r.json()["data"]["hypotheses"]["loyer_plafond"]["PLUS"] == 6.49


@pytest.mark.parametrize("onglet", ["hypotheses", "programme", "prix-revient", "financement", "reservations",
                                    "bilan-operation", "exploitation", "synthese", "complet"])
def test_exports(client, onglet):
    r = client.get(f"/api/simulations/operation/1/export/{onglet}")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert len(wb.sheetnames) == (8 if onglet == "complet" else 1)


def test_bareme_dans_referentiels(client):
    refs = client.get("/api/referentiels").json()
    assert any(r["code"] == "bareme_reglementaire" for r in refs)
