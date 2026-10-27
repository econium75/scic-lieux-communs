"""Suivi des projets et facturation : parcours de la SCIC, documents générés, factures déclenchées par l'avancement."""
from datetime import date, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def _h(client, user, pwd):
    r = client.post("/api/auth/login", data={"username": user, "password": pwd})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module")
def H(client):
    return {"saisie": _h(client, "saisie", "saisie123"), "lecture": _h(client, "lecture", "lecture123"),
            "admin": _h(client, "admin", "admin123")}


TIERS = {"collectivite": {"nom": "Ville d'Orly", "adresse": "7 avenue Adrien-Raynal, 94310 Orly", "contact": "M. le Maire"},
         "promoteur": {"nom": "Promoteur Exemple SAS", "siren": "987654321", "contact": "Mme Martin, directrice des programmes"},
         "financeur": {"nom": "CNAV"}}


@pytest.fixture(scope="module")
def projet(client, H):
    r = client.post("/api/projets", headers=H["saisie"], json={
        "nom": "Habitat participatif Les Tilleuls", "client_nom": "Coopérative Les Tilleuls", "client_adresse": "12 rue des Lilas, Orly",
        "client_siren": "123456789", "client_contact": "Mme Durand, présidente", "ca_previsionnel_ttc": 4_800_000, "taux_tva": 0.2,
        "operation_id": 1, "tiers": TIERS})
    assert r.status_code == 200, r.text
    return r.json()


def _etape(p, nom):
    return next(e for e in p["etapes"] if e["nom"] == nom)


def _texte_pdf(data: bytes) -> str:
    return " ".join(p.extract_text() for p in PdfReader(BytesIO(data)).pages).replace(" ", " ")


def test_projet_cree_avec_le_parcours_de_la_scic(projet):
    p = projet
    assert p["code"].startswith(f"PRJ-{date.today().year}-")
    assert len(p["etapes"]) == 17
    assert [e["phase"] for e in p["etapes"]][::6] == ["Identification et qualification des prospects", "Incubation", "Phase opérationnelle"]
    assert p["etapes"][0]["statut"] == "en_cours" and p["etapes"][1]["statut"] == "a_venir"
    assert p["montant_ht"] == 144_000  # 3 % du chiffre d'affaires prévisionnel TTC
    assert _etape(p, "Diagnostic")["montant_ht"] == 15_000
    assert _etape(p, "Facturation au promoteur")["montant_ht"] == 24_000  # 0,5 % de 4,8 M€
    assert _etape(p, "Esquisse et permis de construire")["montant_ht"] == 14_400  # 10 % des honoraires
    assert p["total_facturable_ht"] == 53_400 and p["part_honoraires_planifiee"] == pytest.approx(0.10)
    assert p["financements"] == [{"etape_id": _etape(p, "Financement de la phase d'incubation")["id"], "nom": "Financement de la phase d'incubation",
                                  "montant": 60_000, "financeur": "CNAV", "obtenu": False}]
    assert p["operation_nom"] == "Orly Ecohome" and p["avancement"] == 0


def test_lecture_seule_pour_le_profil_lecture(client, H, projet):
    assert client.get(f"/api/projets/{projet['id']}", headers=H["lecture"]).status_code == 200
    assert client.post("/api/projets", headers=H["lecture"], json={"nom": "x"}).status_code == 403
    assert client.get("/api/factures", headers=H["lecture"]).status_code == 403
    assert client.get("/api/projets").status_code == 401


def test_documents_generes_avec_les_parties(client, H, projet):
    modeles = {m["titre"]: m for m in client.get("/api/projets/modeles", headers=H["saisie"]).json()["modeles"]}
    assert {"Convention d'objectifs", "Contrat d'honoraires", "Convention d'accompagnement"} <= set(modeles)
    r = client.post(f"/api/projets/{projet['id']}/documents", headers=H["saisie"], json={"modele_id": modeles["Contrat d'honoraires"]["id"]})
    assert r.status_code == 200, r.text
    doc_id = r.json()["document_id"]
    doc = client.get(f"/api/documents/{doc_id}", headers=H["saisie"]).json()
    assert doc["dossier"] == "Projets" and doc["est_pdf"]
    texte = _texte_pdf(client.get(f"/api/documents/{doc_id}/versions/1/fichier", headers=H["saisie"]).content)
    assert "Promoteur Exemple SAS" in texte and "Faire Commun" in texte and "Mme Martin" in texte
    assert "144 000,00" in texte and "4 800 000,00" in texte and "3 %" in texte  # honoraires, CA, taux
    assert "14 400,00" in texte and "Phase opérationnelle" in texte  # échéancier par phase
    assert "champ inconnu" not in texte
    r = client.post(f"/api/projets/{projet['id']}/documents", headers=H["saisie"], json={"modele_id": modeles["Convention d'objectifs"]["id"]})
    texte = _texte_pdf(client.get(f"/api/documents/{r.json()['document_id']}/versions/1/fichier", headers=H["saisie"]).content)
    assert "Ville d'Orly" in texte and "CNAV" in texte and "M. le Maire" in texte


def test_apercu_modele_signale_les_champs_inconnus(client, H):
    r = client.post("/api/projets/modeles/apercu", headers=H["saisie"], json={"contenu": "# Test\nBonjour {{client.nom}} {{nimporte.quoi}}"})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    texte = _texte_pdf(r.content)
    assert "Coopérative d'habitants Les Tilleuls" in texte and "champ inconnu : nimporte.quoi" in texte


def test_diagnostic_termine_facture_la_collectivite(client, H, projet):
    e1 = _etape(projet, "Diagnostic")
    r = client.post(f"/api/projets/{projet['id']}/etapes/{e1['id']}/statut", headers=H["saisie"], json={"statut": "terminee"})
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["etapes"][0]["statut"] == "terminee" and p["etapes"][1]["statut"] == "en_cours"
    f = p["etapes"][0]["facture"]
    assert f["statut"] == "brouillon" and f["total_ttc"] == 18_000
    assert client.get(f"/api/factures/{f['id']}", headers=H["admin"]).json()["client_nom"] == "Ville d'Orly"
    notifs = client.get("/api/notifications", headers=H["admin"]).json()["rows"]
    assert any(n["titre"].startswith("Facture à émettre") for n in notifs)
    # le même événement ne crée pas de second brouillon
    client.post(f"/api/projets/{projet['id']}/etapes/{e1['id']}/statut", headers=H["saisie"], json={"statut": "terminee"})
    assert sum(1 for x in client.get(f"/api/projets/{projet['id']}", headers=H["saisie"]).json()["factures"] if x["etape_id"] == e1["id"]) == 1


def test_financement_cnav_sans_facture(client, H, projet):
    e = _etape(projet, "Financement de la phase d'incubation")
    p = client.post(f"/api/projets/{projet['id']}/etapes/{e['id']}/statut", headers=H["saisie"], json={"statut": "terminee"}).json()
    assert _etape(p, e["nom"])["facture"] is None and p["financements"][0]["obtenu"] is True


def test_emission_numerotee_et_verrouillee(client, H, projet):
    p = client.get(f"/api/projets/{projet['id']}", headers=H["saisie"]).json()
    fid = _etape(p, "Diagnostic")["facture"]["id"]
    assert client.post(f"/api/factures/{fid}/emettre", headers=H["saisie"]).status_code == 403  # réservé à l'administrateur
    assert "BROUILLON" in _texte_pdf(client.get(f"/api/factures/{fid}/pdf", headers=H["saisie"]).content)
    r = client.post(f"/api/factures/{fid}/emettre", headers=H["admin"], json={})
    assert r.status_code == 200, r.text
    f = r.json()
    assert f["numero"] == f"FC-{date.today().year}-0001" and f["statut"] == "emise"
    assert f["date_echeance"] == (date.today() + timedelta(days=30)).isoformat()
    texte = _texte_pdf(client.get(f"/api/factures/{fid}/pdf", headers=H["admin"]).content)
    assert f["numero"] in texte and "Pénalités de retard" in texte and "prestation de services" in texte and "Ville d'Orly" in texte
    assert client.put(f"/api/factures/{fid}", headers=H["admin"], json={"objet": "x"}).status_code == 409
    assert client.delete(f"/api/factures/{fid}", headers=H["admin"]).status_code == 409
    # une étape facturée ne peut plus être rouverte ni changer de règle
    e1 = _etape(p, "Diagnostic")
    assert client.post(f"/api/projets/{projet['id']}/etapes/{e1['id']}/statut", headers=H["saisie"], json={"statut": "en_cours"}).status_code == 409
    etapes = [{**e, "facturation": {**e["facturation"], "montant": 12_000}} if e["id"] == e1["id"] else e for e in p["etapes"]]
    assert client.put(f"/api/projets/{projet['id']}", headers=H["saisie"], json={"etapes": etapes}).status_code == 409
    p = client.get(f"/api/projets/{projet['id']}", headers=H["saisie"]).json()
    assert p["facture_ht"] == 15_000 and p["reste_a_facturer_ht"] == 38_400


def test_numerotation_continue_avoir_et_ca_verrouille(client, H, projet):
    e = _etape(client.get(f"/api/projets/{projet['id']}", headers=H["saisie"]).json(), "Facturation au promoteur")
    p = client.post(f"/api/projets/{projet['id']}/etapes/{e['id']}/statut", headers=H["saisie"], json={"statut": "terminee"}).json()
    fid = _etape(p, e["nom"])["facture"]["id"]
    assert client.get(f"/api/factures/{fid}", headers=H["admin"]).json()["client_nom"] == "Promoteur Exemple SAS"
    f2 = client.post(f"/api/factures/{fid}/emettre", headers=H["admin"], json={}).json()
    assert f2["numero"] == f"FC-{date.today().year}-0002" and f2["total_ht"] == 24_000
    # la facture au promoteur est émise : le chiffre d'affaires prévisionnel ne peut plus changer
    assert client.put(f"/api/projets/{projet['id']}", headers=H["saisie"], json={"ca_previsionnel_ttc": 5_000_000}).status_code == 409
    assert client.get(f"/api/projets/{projet['id']}", headers=H["saisie"]).json()["ca_previsionnel_ttc"] == 4_800_000
    a = client.post(f"/api/factures/{fid}/avoir", headers=H["admin"]).json()
    assert a["type"] == "avoir" and a["total_ttc"] == -f2["total_ttc"]
    assert client.post(f"/api/factures/{fid}/avoir", headers=H["admin"]).status_code == 409
    a = client.post(f"/api/factures/{a['id']}/emettre", headers=H["admin"], json={}).json()
    assert a["numero"] == f"AV-{date.today().year}-0001"
    assert client.get(f"/api/factures/{fid}", headers=H["admin"]).json()["statut"] == "annulee"
    # l'étape redevient facturable : la rouvrir puis la terminer recrée un brouillon
    client.post(f"/api/projets/{projet['id']}/etapes/{e['id']}/statut", headers=H["saisie"], json={"statut": "en_cours"})
    p = client.post(f"/api/projets/{projet['id']}/etapes/{e['id']}/statut", headers=H["saisie"], json={"statut": "terminee"}).json()
    assert _etape(p, e["nom"])["facture"]["statut"] == "brouillon"


def test_paiement_retard_et_indicateurs(client, H, projet):
    p = client.get(f"/api/projets/{projet['id']}", headers=H["saisie"]).json()
    f1 = _etape(p, "Diagnostic")["facture"]["id"]
    assert client.post(f"/api/factures/{f1}/payer", headers=H["admin"], json={"date_paiement": date.today().isoformat()}).json()["statut"] == "payee"
    data = client.get("/api/factures", headers=H["saisie"]).json()
    ind = data["indicateurs"]
    assert ind["encaisse_annee_ttc"] == 18_000 and ind["brouillons"] >= 1
    # date d'émission antérieure à la dernière facture : refusée (ordre chronologique)
    brouillon = next(f for f in data["factures"] if f["statut"] == "brouillon" and f["type"] == "facture")
    hier = (date.today() - timedelta(days=1)).isoformat()
    assert client.post(f"/api/factures/{brouillon['id']}/emettre", headers=H["admin"], json={"date_emission": hier}).status_code == 422


def test_facture_en_retard_notifiee_une_fois(client, H):
    f = client.post("/api/factures", headers=H["admin"], json={"client_nom": "Commune d'Orly", "objet": "Étude",
                                                              "lignes": [{"designation": "Étude", "quantite": 1, "prix_unitaire_ht": 1000, "taux_tva": 0.2}]}).json()
    f = client.post(f"/api/factures/{f['id']}/emettre", headers=H["admin"], json={}).json()
    from app.db import engine
    from app.models import Facture
    from sqlmodel import Session
    with Session(engine) as s:
        x = s.get(Facture, f["id"])
        x.date_echeance = date.today() - timedelta(days=1)
        s.add(x)
        s.commit()
    client.get("/api/factures", headers=H["admin"])
    client.get("/api/factures", headers=H["admin"])
    notifs_admin = [n for n in client.get("/api/notifications?limit=100", headers=H["admin"]).json()["rows"] if f["numero"] in n["titre"]]
    assert len(notifs_admin) == 1
    assert client.get(f"/api/factures/{f['id']}", headers=H["admin"]).json()["en_retard"] is True


def test_journal_des_ventes(client, H):
    r = client.get("/api/factures/export.xlsx", headers=H["saisie"])
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    numeros = [c.value for c in ws["A"][1:]]
    assert numeros[0].startswith("AV-") and any(n.startswith("FC-") for n in numeros)


def test_parcours_lu_avec_des_regles_completes(client, H):
    etapes = client.get("/api/projets/parcours", headers=H["lecture"]).json()["etapes"]
    pc = next(e for e in etapes if e["nom"] == "Esquisse et permis de construire")
    assert pc["facturation"] == {"mode": "pourcentage", "montant": 0.0, "taux": 0.1, "base": "honoraires", "payeur": "promoteur", "nature": "facture"}
    assert next(e for e in etapes if e["nom"].startswith("Financement"))["facturation"]["nature"] == "financement"


def test_parcours_type_modifiable_par_admin(client, H):
    assert client.put("/api/projets/parcours", headers=H["saisie"], json={"etapes": []}).status_code == 403
    r = client.put("/api/projets/parcours", headers=H["admin"], json={"etapes": [
        {"phase": "Amont", "nom": "Convention", "facturation": {"mode": "forfait", "montant": 5000, "payeur": "client"}, "modeles": ["Convention d'accompagnement"]},
        {"phase": "Aval", "nom": "Études", "facturation": {"mode": "pourcentage", "taux": 0.7, "base": "honoraires", "payeur": "client"}}]})
    assert r.status_code == 200 and len(r.json()["etapes"]) == 2
    trop = [{"nom": "A", "facturation": {"mode": "pourcentage", "taux": 0.8, "base": "honoraires"}},
            {"nom": "B", "facturation": {"mode": "pourcentage", "taux": 0.4, "base": "honoraires"}}]
    assert client.put("/api/projets/parcours", headers=H["admin"], json={"etapes": trop}).status_code == 422
    p = client.post("/api/projets", headers=H["saisie"], json={"nom": "Projet 2", "montant_ht": 10000}).json()
    assert [e["nom"] for e in p["etapes"]] == ["Convention", "Études"] and p["etapes"][0]["modeles"]
    assert [e["montant_ht"] for e in p["etapes"]] == [5000, 7000] and p["etapes"][0]["phase"] == "Amont"
    assert p["code"].endswith("-002")


def test_plans_action(client, H):
    plans = client.get("/api/projets/plans-action", headers=H["lecture"]).json()
    assert any(a["titre"] == "Portage des béguinages" for a in plans["actions"])
    assert plans["fonciere"]["annees"][0] == {"annee": 2027, "idf_projets": 1, "idf_logements": 16, "idf_detail": "Habitat partagé (Orly : 16 logements)",
                                              "hors_idf_projets": 2, "hors_idf_logements": 0}
    assert client.put("/api/projets/plans-action", headers=H["lecture"], json=plans).status_code == 403
    plans["actions"][0]["statut"] = "en_cours"
    plans["fonciere"]["cout_moyen_logement"] = 250_000
    r = client.put("/api/projets/plans-action", headers=H["saisie"], json=plans)
    assert r.status_code == 200 and r.json()["actions"][0]["statut"] == "en_cours"
    assert client.put("/api/projets/plans-action", headers=H["saisie"], json={**plans, "actions": [{"titre": "x", "statut": "zz"}]}).status_code == 422


def test_identite_et_referentiels_internes(client, H):
    ident = client.get("/api/projets/identite", headers=H["saisie"]).json()
    assert ident["raison_sociale"] == "Faire Commun"
    assert client.put("/api/projets/identite", headers=H["saisie"], json={"siret": "x"}).status_code == 403
    r = client.put("/api/projets/identite", headers=H["admin"], json={"siret": "123 456 789 00012", "delai_paiement_jours": 45})
    assert r.json()["siret"] == "123 456 789 00012" and r.json()["delai_paiement_jours"] == 45
    codes = {x["code"] for x in client.get("/api/referentiels").json()}
    assert not codes & {"identite_scic", "parcours_projet", "plans_action"}
