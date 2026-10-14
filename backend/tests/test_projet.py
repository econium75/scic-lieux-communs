"""Gestion de projet : planning de Gantt, dépendances, cascade, phases, notifications, export."""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def _h(client, user, pwd):
    r = client.post("/api/auth/login", data={"username": user, "password": pwd})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module")
def H(client):
    return {"saisie": _h(client, "saisie", "saisie123"), "lecture": _h(client, "lecture", "lecture123")}


@pytest.fixture(scope="module")
def op_id(client):
    return next(o["id"] for o in client.get("/api/operations").json() if o["code"] == "ORLY-ECO-2027")


def test_planning_indicatif_orly(client, op_id):
    p = client.get(f"/api/projet/operation/{op_id}").json()
    assert len(p["phases"]) == 8
    assert p["projet"]["nb_taches"] >= 10 and p["projet"]["nb_jalons"] >= 3
    incubation = next(x for x in p["phases"] if x["nom"] == "Incubation")
    assert incubation["planifiee_par_taches"] and incubation["avancement"] == 100
    assert any(t["dependances"] for t in p["taches"])


def test_creation_dependance_cascade_et_notifications(client, H, op_id):
    p = client.get(f"/api/projet/operation/{op_id}").json()
    phase = next(x for x in p["phases"] if x["nom"] == "Livraison")["id"]
    a = client.post("/api/projet/taches", headers=H["saisie"], json={
        "operation_id": op_id, "phase_id": phase, "titre": "Visite de pré-réception", "debut": "2029-03-01",
        "fin": "2029-03-05", "assignes": ["lecture"]}).json()["tache"]
    assert any(n["titre"].startswith("Tâche de planning assignée") for n in
               client.get("/api/notifications", headers=H["lecture"]).json()["rows"])
    # Une tâche créée après un prédécesseur qui finit plus tard est repoussée
    b = client.post("/api/projet/taches", headers=H["saisie"], json={
        "operation_id": op_id, "phase_id": phase, "titre": "Levée des réserves", "debut": "2029-03-02",
        "fin": "2029-03-12", "dependances": [a["id"]], "assignes": ["lecture"]}).json()["tache"]
    assert b["debut"] == "2029-03-06" and b["fin"] == "2029-03-16"  # durée conservée (11 jours)
    # Allonger A repousse B en cascade, et ses assignés sont prévenus
    r = client.put(f"/api/projet/taches/{a['id']}", headers=H["saisie"], json={"fin": "2029-03-10"}).json()
    assert r["decalees"] == [b["id"]]
    t = {x["id"]: x for x in client.get(f"/api/projet/operation/{op_id}").json()["taches"]}
    assert t[b["id"]]["debut"] == "2029-03-11" and t[b["id"]]["fin"] == "2029-03-21"
    assert any("Planning modifié" in n["titre"] for n in client.get("/api/notifications", headers=H["lecture"]).json()["rows"])
    # Sans cascade : le successeur ne bouge pas
    client.put(f"/api/projet/taches/{a['id']}", headers=H["saisie"], json={"fin": "2029-03-15", "cascade": False})
    t = {x["id"]: x for x in client.get(f"/api/projet/operation/{op_id}").json()["taches"]}
    assert t[b["id"]]["debut"] == "2029-03-11"
    # Boucle de dépendances refusée
    assert client.put(f"/api/projet/taches/{a['id']}", headers=H["saisie"], json={"dependances": [b["id"]]}).status_code == 422
    # Terminé → 100 % et le créateur est notifié
    client.put(f"/api/projet/taches/{b['id']}", headers=H["saisie"], json={"statut": "termine"})
    t = {x["id"]: x for x in client.get(f"/api/projet/operation/{op_id}").json()["taches"]}
    assert t[b["id"]]["avancement"] == 100
    # Suppression : la dépendance disparaît chez le successeur
    assert client.delete(f"/api/projet/taches/{a['id']}", headers=H["saisie"]).json()["ok"]
    t = {x["id"]: x for x in client.get(f"/api/projet/operation/{op_id}").json()["taches"]}
    assert a["id"] not in t and t[b["id"]]["dependances"] == []


def test_jalon_et_validations(client, H, op_id):
    j = client.post("/api/projet/taches", headers=H["saisie"], json={
        "operation_id": op_id, "type": "jalon", "titre": "Signature VEFA", "debut": "2027-09-30", "fin": "2027-10-30"}).json()["tache"]
    assert j["fin"] == j["debut"]
    assert client.post("/api/projet/taches", headers=H["saisie"], json={
        "operation_id": op_id, "titre": "x", "debut": "2027-10-10", "fin": "2027-10-01"}).status_code == 422


def test_lecteur_ne_modifie_pas(client, H, op_id):
    assert client.post("/api/projet/taches", headers=H["lecture"], json={
        "operation_id": op_id, "titre": "x", "debut": "2027-10-10"}).status_code == 403


def test_decaler_phase(client, H, op_id):
    p = client.get(f"/api/projet/operation/{op_id}").json()
    ph = next(x for x in p["phases"] if x["nom"] == "Exploitation")
    avant = ph["debut"]
    client.post(f"/api/projet/operation/{op_id}/decaler-phase", headers=H["saisie"], json={"phase_id": ph["id"], "jours": 14})
    apres = next(x for x in client.get(f"/api/projet/operation/{op_id}").json()["phases"] if x["id"] == ph["id"])["debut"]
    from datetime import date
    assert (date.fromisoformat(apres) - date.fromisoformat(avant)).days == 14


def test_phase_sans_tache_avec_dates(client, H, op_id):
    p = client.get(f"/api/projet/operation/{op_id}").json()
    ph = next(x for x in p["phases"] if x["nom"] == "Clôture")
    client.put(f"/api/phases/{ph['id']}", headers=H["saisie"], json={"date_debut": "2030-05-01", "date_cible": "2030-06-30"})
    ph = next(x for x in client.get(f"/api/projet/operation/{op_id}").json()["phases"] if x["id"] == ph["id"])
    assert ph["debut"] == "2030-05-01" and ph["fin"] == "2030-06-30" and not ph["planifiee_par_taches"]


def test_export_excel(client, op_id):
    r = client.get(f"/api/projet/operation/{op_id}/export.xlsx")
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    assert ws["B4"].value == "Incubation"
