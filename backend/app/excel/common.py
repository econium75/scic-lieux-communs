"""Styles openpyxl partagés — chartes graphiques des gabarits SCIC."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Couleurs cohérentes avec la UI (brand-700, paper-200, ok-500, danger-500...)
BRAND = "1F3A68"
PAPER_DARK = "E8E6DC"
PAPER_LIGHT = "F3F2EC"
OK = "2E7D32"
DANGER = "C53030"
WARN = "D97706"
INK_DARK = "1A1F2E"
INK_LIGHT = "6B7390"

THIN = Side(border_style="thin", color="D4D2C8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY_FMT = '#,##0 "€"'
MONEY_FMT_NEG_RED = '#,##0 "€";[Red]-#,##0 "€"'
PCT_FMT = "0.00%"
NUM_FMT = "#,##0"


def header_cell(cell, text: str, bg: str = BRAND, fg: str = "FFFFFF"):
    cell.value = text
    cell.font = Font(bold=True, color=fg)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BOX


def rubrique_cell(cell, text: str):
    cell.value = text
    cell.font = Font(bold=True, color=INK_DARK)
    cell.fill = PatternFill("solid", fgColor=PAPER_DARK)
    cell.alignment = Alignment(horizontal="left")


def total_row_cell(cell, text_or_val, is_number: bool = False, color: str = INK_DARK):
    cell.value = text_or_val
    cell.font = Font(bold=True, color=color)
    cell.fill = PatternFill("solid", fgColor=PAPER_LIGHT)
    if is_number:
        cell.number_format = MONEY_FMT
        cell.alignment = Alignment(horizontal="right")


def set_column_widths(ws, widths: dict[str, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
