"""Export Excel de la commercialisation BRS : hypothèses, ménages candidats, programme et subventions."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .common import BRAND, DANGER, INK_LIGHT, MONEY_FMT, OK, WARN, header_cell, rubrique_cell, set_column_widths
from ..engines.commercialisation import LIBELLE_TRANCHE, TRANCHES

PCT = "0.0%"
SAISIE = PatternFill("solid", fgColor="FFF6CC")
COULEUR_VERDICT = {"financable": OK, "a_retravailler": WARN, "non_eligible": DANGER, "a_completer": INK_LIGHT}


def _titre(ws, titre: str, sous_titre: str) -> None:
    ws["A1"] = titre
    ws["A1"].font = Font(bold=True, size=14, color=BRAND)
    ws["A2"] = sous_titre
    ws["A2"].font = Font(italic=True, size=9, color=INK_LIGHT)


def _feuille_hypotheses(ws, h: dict, nom_operation: str) -> None:
    _titre(ws, f"Simulation BRS — Hypothèses · {nom_operation}", "Cellules jaunes = valeurs saisies dans Téo.")
    for c, t in enumerate(["Typologie", "Surface habitable (m²)", "Prix de vente BRS TTC", "Nb logements",
                           "Subv. déciles 3-4", "Subv. déciles 5-6", "Subv. déciles 7-9"], 1):
        header_cell(ws.cell(row=4, column=c), t)
    r = 5
    for t in h["typologies"]:
        vals = [t.get("code"), t.get("surface"), t.get("prix_ttc"), t.get("nb"),
                *[(t.get("subventions") or {}).get(k) for k in TRANCHES]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c > 1:
                cell.fill = SAISIE
            if c in (3, 5, 6, 7):
                cell.number_format = MONEY_FMT
        r += 1

    r += 1
    rubrique_cell(ws.cell(row=r, column=1), "Paramètres financiers")
    r += 1
    for libelle, cle, fmt in [
        ("Redevance foncière (€/m² SH/mois)", "redevance_m2_mois", '0.00 "€"'),
        ("Charges de copropriété (€/m² SH/mois)", "charges_m2_mois", '0.00 "€"'),
        ("Taxe foncière annuelle estimée (€)", "taxe_fonciere_an", MONEY_FMT),
        ("Taux du prêt (hors assurance)", "taux_pret", "0.00%"),
        ("Taux assurance emprunteur (annuel, sur capital)", "taux_assurance", "0.00%"),
        ("Durée du prêt (années)", "duree_ans", "0"),
        ("Frais d'acquisition (% du prix)", "frais_acquisition", PCT),
        ("Taux d'effort maximal (HCSF)", "taux_effort_max", PCT),
        ("Reste à vivre minimum par UC (€/mois)", "reste_a_vivre_min_uc", MONEY_FMT),
    ]:
        ws.cell(row=r, column=1, value=libelle)
        cell = ws.cell(row=r, column=2, value=h.get(cle))
        cell.number_format, cell.fill = fmt, SAISIE
        r += 1

    r += 1
    pl = h.get("plafonds") or {}
    rubrique_cell(ws.cell(row=r, column=1), f"Plafonds de ressources BRS — revenu fiscal de référence N-2 · {pl.get('zone', '')}")
    r += 1
    ws.cell(row=r, column=1, value=pl.get("source", "")).font = Font(italic=True, size=9, color=INK_LIGHT)
    r += 1
    for i, libelle in enumerate(["1", "2", "3", "4", "5 et +"]):
        ws.cell(row=r, column=1, value=f"{libelle} personne{'s' if i else ''}")
        montants = pl.get("montants") or []
        cell = ws.cell(row=r, column=2, value=montants[i] if i < len(montants) else None)
        cell.number_format, cell.fill = MONEY_FMT, SAISIE
        r += 1

    r += 1
    d = h.get("deciles") or {}
    rubrique_cell(ws.cell(row=r, column=1), f"Déciles de revenu disponible par UC ({d.get('source', '')})")
    r += 1
    for k in ("D3", "D4", "D5", "D6", "D7", "D8", "D9"):
        ws.cell(row=r, column=1, value=k + (" (médiane)" if k == "D5" else ""))
        cell = ws.cell(row=r, column=2, value=d.get(k))
        cell.number_format, cell.fill = MONEY_FMT, SAISIE
        r += 1
    set_column_widths(ws, {"A": 48, "B": 22, "C": 22, "D": 14, "E": 18, "F": 18, "G": 18})


def _feuille_menages(ws, resultats: list[dict]) -> None:
    _titre(ws, "Ménages candidats", "Éligibilité, tranche de subvention, coût mensuel et solvabilité de chaque ménage.")
    cols = [
        ("Réf.", "reference", None), ("Typo", "typologie", None), ("Personnes", "personnes", "0"), ("UC", "uc", "0.0"),
        ("Plafond", "plafond", MONEY_FMT), ("Éligible", None, None), ("Revenu dispo / UC", "revenu_par_uc", MONEY_FMT),
        ("Tranche", "tranche_libelle", None), ("Prix BRS", "prix", MONEY_FMT), ("Subvention", "subvention", MONEY_FMT),
        ("Prix net", "prix_net", MONEY_FMT), ("Frais", "frais", MONEY_FMT), ("Emprunt", "emprunt", MONEY_FMT),
        ("Mensualité", "mensualite", MONEY_FMT), ("Assurance", "assurance", MONEY_FMT), ("Redevance", "redevance", MONEY_FMT),
        ("Charges", "charges", MONEY_FMT), ("Taxe foncière", "taxe_fonciere", MONEY_FMT), ("Coût mensuel", "cout_mensuel", MONEY_FMT),
        ("Taux d'effort", "taux_effort", PCT), ("Reste à vivre", "reste_a_vivre", MONEY_FMT),
        ("Reste à vivre / UC", "reste_a_vivre_uc", MONEY_FMT), ("Verdict", "verdict_libelle", None), ("Motifs", None, None),
    ]
    for c, (t, _, _) in enumerate(cols, 1):
        header_cell(ws.cell(row=4, column=c), t)
    for r, q in enumerate(resultats, 5):
        for c, (t, cle, fmt) in enumerate(cols, 1):
            if t == "Éligible":
                v = "Plafond à saisir" if q["eligible"] is None else ("OUI" if q["eligible"] else "NON")
            elif t == "Motifs":
                v = " · ".join(q["motifs"])
            else:
                v = q.get(cle)
            cell = ws.cell(row=r, column=c, value=v)
            if fmt:
                cell.number_format = fmt
            if t == "Verdict":
                cell.font = Font(bold=True, color=COULEUR_VERDICT.get(q["verdict"], BRAND))
    set_column_widths(ws, {chr(64 + i): 14 for i in range(1, 24)} | {"H": 16, "W": 34, "X": 60})


def _feuille_programme(ws, prog: dict) -> None:
    _titre(ws, "Programme BRS — répartition des subventions par déciles", "Nombre de logements de chaque typologie par tranche de déciles.")
    for c, t in enumerate(["Typologie", "Nb lgts", *[LIBELLE_TRANCHE[k] for k in TRANCHES], "Contrôle répartition",
                           "Subv. totale", "Prix brut total", "Prix net total"], 1):
        header_cell(ws.cell(row=3, column=c), t)
    r = 4
    for l in prog["lignes"]:
        vals = [l["code"], l["nb"], *[l["repartition"][k] for k in TRANCHES], l["controle"], l["subvention"], l["prix_brut"], l["prix_net"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c >= 7:
                cell.number_format = MONEY_FMT
            if c == 6:
                cell.font = Font(bold=True, color=OK if l["ecart"] == 0 else DANGER)
        r += 1
    tot = prog["total"]
    for c, v in enumerate(["Total", tot["nb"], *[tot[k] for k in TRANCHES], "", tot["subvention"], tot["prix_brut"], tot["prix_net"]], 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(bold=True)
        if c >= 7:
            cell.number_format = MONEY_FMT
    r += 1
    for libelle, valeurs in (("Quotité réelle", prog["quotites_reelles"]), ("Quotité cible", prog["quotites_cibles"])):
        ws.cell(row=r, column=1, value=libelle)
        for i, k in enumerate(TRANCHES):
            ws.cell(row=r, column=3 + i, value=valeurs[k]).number_format = PCT
        r += 1
    ws.cell(row=r, column=1, value="Subvention moyenne par logement")
    ws.cell(row=r, column=7, value=prog["subvention_moyenne"]).number_format = MONEY_FMT
    for row in ws.iter_rows(min_row=3, max_row=r):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
    set_column_widths(ws, {"A": 30, "B": 10, "C": 14, "D": 14, "E": 14, "F": 20, "G": 16, "H": 18, "I": 18})


def build_commercialisation(h: dict, resultats: list[dict], prog: dict, nom_operation: str = "") -> BytesIO:
    wb = Workbook()
    _feuille_hypotheses(wb.active, h, nom_operation)
    wb.active.title = "Hypothèses"
    _feuille_menages(wb.create_sheet("Ménages candidats"), resultats)
    _feuille_programme(wb.create_sheet("Programme & subventions"), prog)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
