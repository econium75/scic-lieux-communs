"""Exports Excel du simulateur — un onglet par écran, ou le classeur complet."""
from io import BytesIO
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from .common import (
    BRAND, PAPER_LIGHT, INK_LIGHT, MONEY_FMT, MONEY_FMT_NEG_RED, PCT_FMT, NUM_FMT, BOX,
    header_cell, rubrique_cell,
)
from ..engines.simulateur import PRODUITS, RUBRIQUES

DEC_FMT = '#,##0.00'
P5 = [*PRODUITS, "Total"]


def _titre(ws, titre: str, sous_titre: str = "") -> int:
    ws["A1"] = titre
    ws["A1"].font = Font(bold=True, size=14, color=BRAND)
    if sous_titre:
        ws["A2"] = sous_titre
        ws["A2"].font = Font(italic=True, color=INK_LIGHT)
    return 4


def _entetes(ws, row: int, labels: list[str], col: int = 1) -> int:
    for i, l in enumerate(labels):
        header_cell(ws.cell(row=row, column=col + i), l)
    return row + 1


def _ligne(ws, row: int, label: str, valeurs: list[Any], fmt: str = MONEY_FMT, bold: bool = False) -> int:
    c = ws.cell(row=row, column=1, value=label)
    c.border = BOX
    if bold:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=PAPER_LIGHT)
    for i, v in enumerate(valeurs, start=2):
        cell = ws.cell(row=row, column=i, value=v)
        cell.border = BOX
        if isinstance(v, (int, float)):
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="right")
        if bold:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=PAPER_LIGHT)
    return row + 1


def _par_produit(d: dict, total: bool = True) -> list[float]:
    vals = [d.get(p, 0.0) for p in PRODUITS]
    if total:
        vals.append(d.get("total", sum(vals)))
    return vals


def _largeurs(ws, premiere: int = 46, autres: int = 15, n: int = 12) -> None:
    ws.column_dimensions["A"].width = premiere
    for i in range(2, n + 2):
        ws.column_dimensions[get_column_letter(i)].width = autres


# ─────────────────────────────────────────────────────── 1. hypothèses
def feuille_hypotheses(ws, data: dict, res: dict) -> None:
    h = data["hypotheses"]
    r = _titre(ws, "Hypothèses d'opération", f"{h.get('nom_operation', '')} — {h.get('commune', '')}")

    def bloc(titre: str, items: list[tuple[str, Any, str]]):
        nonlocal r
        rubrique_cell(ws.cell(row=r, column=1), titre)
        r += 1
        for label, v, fmt in items:
            r = _ligne(ws, r, label, [v], fmt)
        r += 1

    bloc("IDENTIFICATION", [
        ("Nom de l'opération", h["nom_operation"], "@"), ("Commune", h["commune"], "@"),
        ("Zone loyers PLAI/PLUS (1bis/1/2/3)", h["zone_123"], "@"), ("Zone A/B/C", h["zone_abc"], "@"),
        ("Année de livraison prévisionnelle", h["annee_livraison"], "0"), ("Millésime des barèmes", h["annee_bareme"], "0"),
    ])
    bloc("COÛTS (HT) — pré-remplissage du prix de revient", [
        ("Charge foncière (€/m² SHAB)", h["charge_fonciere_m2"], DEC_FMT),
        ("Travaux bâtiment + VRD (€/m² SHAB)", h["travaux_m2"], DEC_FMT),
        ("Honoraires (% travaux)", h["honoraires_pct"], PCT_FMT),
        ("Aléas / imprévus (% travaux)", h["aleas_pct"], PCT_FMT),
        ("Frais annexes (€/logement)", h["frais_annexes_logt"], MONEY_FMT),
        ("Décote foncière BRS", h["decote_fonciere_brs_pct"], PCT_FMT),
        *[(f"TVA non récupérable — {p}", h["tva"][p], PCT_FMT) for p in PRODUITS],
        ("TVA sur cession BRS", h["tva_cession_brs"], PCT_FMT),
    ])
    bloc("FINANCEMENT", [
        *[(f"Subvention État {p} (€/logement)", h["subv_etat_logt"][p], MONEY_FMT) for p in PRODUITS],
        ("Subvention collectivités (€/logement)", h["subv_collectivites_logt"], MONEY_FMT),
        ("Subvention Action Logement (€/logement)", h["subv_action_logement_logt"], MONEY_FMT),
        ("Fonds propres cible (% PR TTC)", h["fonds_propres_pct"], PCT_FMT),
        *[x for p in PRODUITS for x in (
            (f"Taux prêt {p}", h["prets"][p]["taux"], PCT_FMT),
            (f"Durée prêt {p} — construction (ans)", h["prets"][p]["duree"], "0"),
            (f"Durée prêt {p} — foncier (ans)", h["prets"][p]["duree_foncier"], "0"),
        )],
        ("Quotité foncière", "Prorata de la charge foncière" if h["quotite_fonciere_mode"] == "prorata"
         else f"{h['quotite_fonciere_pct']:.0%} de la charge foncière HT", "@"),
        ("Prêt CNAV — taux", h["pret_cnav"]["taux"], PCT_FMT), ("Prêt CNAV — durée (ans)", h["pret_cnav"]["duree"], "0"),
        ("Prêt Action Logement — taux", h["pret_action_logement"]["taux"], PCT_FMT),
        ("Prêt Action Logement — durée (ans)", h["pret_action_logement"]["duree"], "0"),
        ("Prêt complémentaire — taux", h["pret_complementaire"]["taux"], PCT_FMT),
        ("Prêt complémentaire — durée (ans)", h["pret_complementaire"]["duree"], "0"),
        ("Prix de vente BRS saisi (€ HT/m² SU)", h["prix_vente_brs_m2_su"], MONEY_FMT),
    ])
    bloc("LOYERS / REDEVANCE ET EXPLOITATION", [
        *[(f"Loyer plafond {p} (€/m² SU/mois)" if p != "BRS" else "Redevance BRS (€/m² SU/mois)",
           h["loyer_plafond"][p], DEC_FMT) for p in PRODUITS],
        ("Indexation annuelle des loyers (IRL)", h["indexation_loyers"], PCT_FMT),
        ("Vacance locative", h["vacance"], PCT_FMT), ("Impayés nets", h["impayes"], PCT_FMT),
        ("Frais de gestion (% loyers encaissés)", h["frais_gestion_pct"], PCT_FMT),
        ("Entretien courant (€/logt/an)", h["entretien_logt"], MONEY_FMT),
        ("Provision gros entretien (€/logt/an)", h["pge_logt"], MONEY_FMT),
        ("Assurances (€/logt/an)", h["assurance_logt"], MONEY_FMT),
        ("Inflation des charges", h["inflation_charges"], PCT_FMT),
        ("TFPB plein tarif (€/logt/an)", h["tfpb_logt"], MONEY_FMT),
        ("Exonération TFPB PLAI/PLUS/PLS (ans)", h["exoneration_tfpb_ans"], "0"),
        ("Parkings (nombre)", h["parkings_nb"], "0"), ("Loyer parking (€/mois)", h["parking_loyer_mensuel"], MONEY_FMT),
        ("RLS (€/logt/mois)", h["rls_logt_mois"], MONEY_FMT), ("Part de logements concernés par la RLS", h["rls_part_logements"], PCT_FMT),
        ("Taux d'actualisation (VAN)", h["taux_actualisation"], PCT_FMT),
        ("Durée de projection (ans)", h["duree_projection"], "0"),
        ("Plafond de surface annexe (m²/logt)", h["annexe_plafond_m2"], DEC_FMT),
        ("Pondération des surfaces annexes", h["annexe_ponderation"], PCT_FMT),
    ])
    _largeurs(ws, 52, 22, 1)


# ─────────────────────────────────────────────────────── 2. programme & loyers
def feuille_programme(ws, data: dict, res: dict) -> None:
    r = _titre(ws, "Programme et calcul des loyers",
               "SU = SHAB + MIN(surface annexe, plafond) × pondération ; loyer = SU × loyer plafond (€/m² SU/mois)")
    pp = res["programme"]["par_produit"]
    r = _entetes(ws, r, ["Produit", "Nb logts", "SHAB (m²)", "SU (m²)", "Loyer €/m² SU/mois", "Loyer annuel (€)"])
    for p in PRODUITS:
        r = _ligne(ws, r, p, [pp[p]["nb"], pp[p]["shab"], pp[p]["su"], pp[p]["loyer_m2"], pp[p]["loyer_annuel"]], DEC_FMT)
    t = res["programme"]["total"]
    r = _ligne(ws, r, "TOTAL", [t["nb"], t["shab"], t["su"], None, t["loyer_annuel"]], DEC_FMT, bold=True)
    r += 1
    for p in PRODUITS:
        if not pp[p]["nb"]:
            continue
        rubrique_cell(ws.cell(row=r, column=1), f"{p}" + (" — redevance foncière" if p == "BRS" else ""))
        r += 1
        r = _entetes(ws, r, ["Typologie", "Nb logts", "SHAB moy. (m²)", "Annexe brute (m²)", "Annexe retenue (m²)",
                             "SU moy. (m²)", "€/m² SU/mois", "Mensuel / logt (€)", "Annuel / logt (€)", "Annuel total (€)"])
        for l in pp[p]["lignes"]:
            if not l["nb"]:
                continue
            r = _ligne(ws, r, l["typologie"], [l["nb"], l["shab_moy"], l["annexe_brute"], l["annexe_retenue"], l["su_moy"],
                                                l["loyer_m2"], l["loyer_mensuel_logt"], l["loyer_annuel_logt"],
                                                l["loyer_annuel_total"]], DEC_FMT)
        r = _ligne(ws, r, f"Total {p}", [pp[p]["nb"], None, None, None, None, None, None, None, pp[p]["loyer_annuel"]],
                   DEC_FMT, bold=True)
        r += 1
    _largeurs(ws, 22, 16, 10)


# ─────────────────────────────────────────────────────── 3. prix de revient
def feuille_prix_revient(ws, data: dict, res: dict) -> None:
    r = _titre(ws, "Prix de revient détaillé (€ HT)", data["hypotheses"].get("nom_operation", ""))
    pr = res["prix_revient"]
    r = _entetes(ws, r, ["Poste", *P5])
    for rub in RUBRIQUES:
        r = _ligne(ws, r, rub, _par_produit(pr["par_rubrique"][rub]), bold=True)
        for l in pr["lignes"]:
            if l["rubrique"] != rub:
                continue
            lib = f"   {l['libelle']}" + ("  (ligne ajoutée)" if l.get("custom") else "")
            r = _ligne(ws, r, lib, [*[l["montants"][p] for p in PRODUITS], l["total_ht"]])
    r += 1
    r = _ligne(ws, r, "PRIX DE REVIENT HT", _par_produit(pr["ht"]), bold=True)
    r = _ligne(ws, r, "Taux de TVA non récupérable", [pr["taux_tva"][p] for p in PRODUITS], PCT_FMT)
    r = _ligne(ws, r, "TVA", _par_produit(pr["tva"]))
    r = _ligne(ws, r, "PRIX DE REVIENT TTC", _par_produit(pr["ttc"]), bold=True)
    r = _ligne(ws, r, "Prix de revient TTC / logement", _par_produit(pr["par_logement"]))
    _ligne(ws, r, "Prix de revient TTC / m² SHAB", _par_produit(pr["par_m2_shab"]))
    _largeurs(ws, 48, 15, 6)


# ─────────────────────────────────────────────────────── 4. plan de financement
def feuille_financement(ws, data: dict, res: dict) -> None:
    r = _titre(ws, "Plan de financement détaillé", "Emprunt principal calculé par équilibre ; montants en €")
    f = res["financement"]
    pr = res["prix_revient"]
    r = _entetes(ws, r, ["", *P5])
    r = _ligne(ws, r, "PRIX DE REVIENT TTC (emplois)", _par_produit(pr["ttc"]), bold=True)
    r += 1
    rubrique_cell(ws.cell(row=r, column=1), "RESSOURCES")
    r += 1
    for l in f["lignes"]:
        lib = l["libelle"] + ("  (complémentaire)" if l["complementaire"] else "")
        r = _ligne(ws, r, lib, [*[l["montants"][p] for p in PRODUITS], l["total"]])
    r = _ligne(ws, r, "Cession affectée (BRS)", _par_produit(f["cession_affectee"]))
    r = _ligne(ws, r, "EMPRUNT D'ÉQUILIBRE (prêt principal)", _par_produit(f["emprunt_equilibre"]), bold=True)
    r = _ligne(ws, r, "   dont quotité foncière", _par_produit(f["quotite_fonciere"]))
    r = _ligne(ws, r, "   dont quotité construction", _par_produit(f["quotite_construction"]))
    r = _ligne(ws, r, "TOTAL RESSOURCES", _par_produit(f["total_ressources"]), bold=True)
    r = _ligne(ws, r, "Contrôle : ressources − prix de revient", _par_produit(f["controle"]), MONEY_FMT_NEG_RED)
    r = _ligne(ws, r, "Excédent de ressources (sur-financement)", _par_produit(f["excedent"]), MONEY_FMT_NEG_RED)
    r += 1
    rubrique_cell(ws.cell(row=r, column=1), "CESSION DES DROITS RÉELS (BRS)")
    r += 1
    c = f["cession_brs"]
    for lib, v, fmt in (("Prix de vente saisi (€ HT/m² SU)", c["prix_saisi_m2"], MONEY_FMT),
                        ("Prix plafond PSLA/BRS de la zone", c["prix_plafond_m2"], MONEY_FMT),
                        ("Prix retenu (€ HT/m² SU)", c["prix_retenu_m2"], MONEY_FMT),
                        ("Surface utile BRS (m²)", c["su"], DEC_FMT),
                        ("Produit de cession HT", c["produit_ht"], MONEY_FMT),
                        ("Produit de cession TTC", c["produit_ttc"], MONEY_FMT),
                        ("Cession affectée au coût du bâti", c["affectee"], MONEY_FMT),
                        ("Marge de cession", c["marge"], MONEY_FMT)):
        r = _ligne(ws, r, lib, [v], fmt)
    r += 1
    rubrique_cell(ws.cell(row=r, column=1), "STRUCTURE (% du prix de revient TTC)")
    r += 1
    for lib, k in (("Subventions", "subventions"), ("Fonds propres", "fonds_propres"),
                   ("Emprunts", "emprunts"), ("Cession BRS", "cession")):
        r = _ligne(ws, r, lib, _par_produit(f["structure"][k]), PCT_FMT)
    r += 1
    rubrique_cell(ws.cell(row=r, column=1), "PRÊTS ET ANNUITÉS (annuités constantes)")
    r += 1
    r = _entetes(ws, r, ["Prêt", "Produit", "Montant", "Taux", "Durée (ans)", "Annuité"])
    for x in f["prets"]:
        ws.cell(row=r, column=1, value=x["libelle"]).border = BOX
        ws.cell(row=r, column=2, value=x["produit"]).border = BOX
        for col, v, fmt in ((3, x["montant"], MONEY_FMT), (4, x["taux"], PCT_FMT), (5, x["duree"], "0"), (6, x["annuite"], MONEY_FMT)):
            cell = ws.cell(row=r, column=col, value=v)
            cell.number_format = fmt
            cell.border = BOX
        r += 1
    _ligne(ws, r, "Total annuités année 1", [None, None, None, None, f["annuites_annee1"]["total"]], bold=True)
    _largeurs(ws, 48, 15, 6)


# ─────────────────────────────────────────────────────── 5. réservations
def feuille_reservations(ws, data: dict, res: dict) -> None:
    r = _titre(ws, "Tableau de réservation des logements", data["hypotheses"].get("nom_operation", ""))
    cols = [("logement", "Logement"), ("typologie", "Typologie"), ("produit", "Produit"), ("categorie", "Catégorie"),
            ("reservataire", "Réservataire"), ("convention", "N° convention"), ("date_signature", "Signature"),
            ("date_echeance", "Échéance"), ("mode", "Gestion"), ("statut", "Statut"), ("contact", "Contact"),
            ("observations", "Observations")]
    r = _entetes(ws, r, [c[1] for c in cols])
    for row in data["reservations"]:
        for i, (k, _) in enumerate(cols, start=1):
            ws.cell(row=r, column=i, value=row.get(k, "")).border = BOX
        r += 1
    r += 1
    s = res["reservations"]
    rubrique_cell(ws.cell(row=r, column=1), "RÉPARTITION ET REPÈRES LÉGAUX")
    r += 1
    r = _entetes(ws, r, ["Catégorie", "Logements", "Part du programme"])
    for c in s["par_categorie"]:
        r = _ligne(ws, r, c["categorie"], [c["nb"], c["part"]], PCT_FMT)
    r += 1
    r = _entetes(ws, r, ["Repère", "Logements", "Part", "Plafond", "Statut"])
    for q in s["quotas"]:
        r = _ligne(ws, r, q["libelle"], [q["nb"], q["part"], q["plafond"], "Dépassé" if q["depasse"] else "OK"], PCT_FMT)
    widths = [14, 11, 9, 28, 30, 18, 12, 12, 10, 14, 22, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ─────────────────────────────────────────────────────── 6. bilan d'opération
def feuille_bilan_operation(ws, data: dict, res: dict) -> None:
    r = _titre(ws, "Bilan d'opération", "Récapitulatif du prix de revient et du plan de financement")
    b = res["bilan_operation"]
    r = _entetes(ws, r, ["", *P5])
    rubrique_cell(ws.cell(row=r, column=1), "EMPLOIS — PRIX DE REVIENT")
    r += 1
    for rub in RUBRIQUES:
        r = _ligne(ws, r, rub, _par_produit(b["emplois"][rub]))
    r = _ligne(ws, r, "SOUS-TOTAL HT", _par_produit(b["ht"]), bold=True)
    r = _ligne(ws, r, "TVA", _par_produit(b["tva"]))
    r = _ligne(ws, r, "TOTAL PRIX DE REVIENT TTC", _par_produit(b["ttc"]), bold=True)
    r = _ligne(ws, r, "Prix de revient / logement", _par_produit(b["par_logement"]))
    r = _ligne(ws, r, "Prix de revient / m² SHAB", _par_produit(b["par_m2_shab"]))
    r += 1
    rubrique_cell(ws.cell(row=r, column=1), "RESSOURCES — FINANCEMENT")
    r += 1
    for l in b["ressources"]:
        r = _ligne(ws, r, l["libelle"], [*[l["montants"][p] for p in PRODUITS], l["total"]])
    r = _ligne(ws, r, "Cession affectée (BRS)", _par_produit(b["cession_affectee"]))
    r = _ligne(ws, r, "Emprunt d'équilibre", _par_produit(b["emprunt_equilibre"]))
    r = _ligne(ws, r, "TOTAL RESSOURCES", _par_produit(b["total_ressources"]), bold=True)
    r = _ligne(ws, r, "Écart ressources − emplois (contrôle)", _par_produit(b["controle"]), MONEY_FMT_NEG_RED)
    r += 1
    rubrique_cell(ws.cell(row=r, column=1), "STRUCTURE DU FINANCEMENT (% du PR TTC)")
    r += 1
    for lib, k in (("Subventions", "subventions"), ("Fonds propres", "fonds_propres"),
                   ("Emprunts", "emprunts"), ("Cession BRS", "cession")):
        r = _ligne(ws, r, lib, _par_produit(b["structure"][k]), PCT_FMT)
    _largeurs(ws, 44, 15, 6)


# ─────────────────────────────────────────────────────── 7. bilan d'exploitation
EXP_COLS = [
    ("annee", "Année n", "0"), ("annee_civile", "Année civile", "0"), ("coef_indexation", "Coef. index.", "0.0000"),
    ("loyers", "Loyers logements", MONEY_FMT), ("redevances_brs", "Redevances BRS (OFS)", MONEY_FMT),
    ("parkings", "Parkings", MONEY_FMT), ("vacance", "Vacance", MONEY_FMT),
    ("impayes", "Impayés", MONEY_FMT), ("rls", "RLS", MONEY_FMT), ("produits_nets", "Produits nets", MONEY_FMT),
    ("frais_gestion", "Frais de gestion", MONEY_FMT), ("entretien", "Entretien courant", MONEY_FMT),
    ("pge", "Provision gros entretien", MONEY_FMT), ("assurance", "Assurance", MONEY_FMT), ("tfpb", "TFPB", MONEY_FMT),
    ("charges", "Total charges", MONEY_FMT), ("autofinancement_brut", "Autofinancement brut", MONEY_FMT),
    ("annuites", "Annuités", MONEY_FMT), ("autofinancement_net", "Autofinancement net", MONEY_FMT_NEG_RED),
    ("cumul", "Cumul", MONEY_FMT_NEG_RED), ("autof_sur_loyers", "Autof. / loyers", PCT_FMT),
    ("annuites_sur_loyers", "Annuités / loyers", PCT_FMT), ("dscr", "DSCR", "0.00"),
]


def feuille_exploitation(ws, data: dict, res: dict, avec_brs: bool = False) -> None:
    e = res["exploitation_avec_brs" if avec_brs else "exploitation"]
    k = e["kpis"]
    titre = ("Bilan d'exploitation prévisionnel — avec BRS (locatif + compte OFS consolidés)" if avec_brs
             else "Bilan d'exploitation prévisionnel — sans BRS (locatif : PLAI, PLUS, PLS, PLI)")
    r = _titre(ws, titre,
               "Montants en € courants ; charges en négatif. Alimenté par les hypothèses, les loyers, le prix de revient et le plan de financement.")
    kp = [
        ("Autofinancement net année 1", k["autofinancement_net_a1"], MONEY_FMT_NEG_RED),
        ("Autofinancement net / loyers année 1", k["autof_sur_loyers_a1"], PCT_FMT),
        ("Annuités / loyers année 1", k["annuites_sur_loyers_a1"], PCT_FMT),
        ("DSCR année 1", k["dscr_a1"], "0.00"),
        ("Autofinancement cumulé à 40 ans", k["cumul_40"], MONEY_FMT_NEG_RED),
        (f"Autofinancement cumulé à {k['nb_annees']} ans", k["cumul_fin"], MONEY_FMT_NEG_RED),
        ("Année de retour à l'équilibre (cumul > 0)", k["annee_equilibre"] or f"> {k['nb_annees']} ans", "0"),
        ("Année de retour des fonds propres", k["annee_retour_fonds_propres"] or "Non atteint", "0"),
        ("TRI fonds propres à 40 ans", k["tri_fp_40"] if k["tri_fp_40"] is not None else "Non atteint", PCT_FMT),
        ("VAN à 40 ans", k["van_40"], MONEY_FMT_NEG_RED),
    ]
    for lib, v, fmt in kp:
        r = _ligne(ws, r, lib, [v], fmt)
    r += 1
    cols_exp = [c for c in EXP_COLS if avec_brs or c[0] != "redevances_brs"]
    r = _entetes(ws, r, [c[1] for c in cols_exp])
    for l in e["lignes"]:
        for i, (key, _, fmt) in enumerate(cols_exp, start=1):
            cell = ws.cell(row=r, column=i, value=l.get(key))
            cell.number_format = fmt
            cell.border = BOX
        r += 1
    if e["brs"] and not avec_brs:
        r += 1
        rubrique_cell(ws.cell(row=r, column=1), "COMPTE OFS — BRS (redevances foncières / portage du foncier)")
        r += 1
        cols = [("annee", "Année n", "0"), ("annee_civile", "Année civile", "0"), ("redevances", "Redevances", MONEY_FMT),
                ("impayes", "Impayés", MONEY_FMT), ("frais_gestion", "Frais de gestion", MONEY_FMT),
                ("annuites", "Annuités prêt foncier", MONEY_FMT), ("solde", "Solde", MONEY_FMT_NEG_RED),
                ("cumul", "Cumul", MONEY_FMT_NEG_RED)]
        r = _entetes(ws, r, [c[1] for c in cols])
        for l in e["brs"]:
            for i, (key, _, fmt) in enumerate(cols, start=1):
                cell = ws.cell(row=r, column=i, value=l.get(key))
                cell.number_format = fmt
                cell.border = BOX
            r += 1
    ws.column_dimensions["A"].width = 40
    for i in range(2, len(EXP_COLS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.freeze_panes = None


# ─────────────────────────────────────────────────────── 8. synthèse
def feuille_synthese(ws, data: dict, res: dict, avec_brs: bool = True) -> None:
    s = res["synthese" if avec_brs else "synthese_sans_brs"]
    r = _titre(ws, "Synthèse — Simulation financière de l'opération" + (" (avec BRS)" if avec_brs else " (sans BRS)"),
               f"{s['titre']} (zones {s['zones']['zone_123']} / {s['zones']['zone_abc']}) — livraison {s['annee_livraison']}")
    for lib, v, fmt in (
        ("Nombre de logements", s["nb_logements"], NUM_FMT), ("SHAB totale (m²)", s["shab"], DEC_FMT),
        ("Prix de revient total TTC", s["prix_revient_ttc"], MONEY_FMT),
        ("Prix de revient / logement", s["prix_revient_logt"], MONEY_FMT),
        ("Prix de revient / m² SHAB", s["prix_revient_m2"], MONEY_FMT),
        ("Autofinancement net — année 1", s["autofinancement_net_a1"], MONEY_FMT_NEG_RED),
        ("DSCR — année 1", s["dscr_a1"], "0.00"),
        ("Année de retour à l'équilibre (cumul > 0)", s["annee_equilibre"] or "Non atteint", "0"),
        ("Fonds propres (total)", s["fonds_propres_total"], MONEY_FMT),
        ("Total emprunts (tous prêts)", s["emprunts_total"], MONEY_FMT),
        ("Annuités année 1", s["annuites_a1"], MONEY_FMT),
    ):
        r = _ligne(ws, r, lib, [v], fmt)
    r += 1
    r = _entetes(ws, r, ["Financement", "Nb logts", "SHAB (m²)", "Prix de revient TTC", "Loyer/redevance (€/m²/mois)", "Loyer annuel (€)"])
    for rep in s["repartition"]:
        r = _ligne(ws, r, rep["produit"], [rep["nb"], rep["shab"], rep["prix_revient_ttc"], rep["loyer_m2"], rep["loyer_annuel"]], MONEY_FMT)
    r += 1
    r = _entetes(ws, r, ["Structure du plan de financement", "Montant (€)", "% du total"])
    for st in s["structure"]:
        r = _ligne(ws, r, st["poste"], [st["montant"], st["part"]])
        ws.cell(row=r - 1, column=3).number_format = PCT_FMT
    r += 1
    if res["alertes"]:
        rubrique_cell(ws.cell(row=r, column=1), "ALERTES")
        r += 1
        for a in res["alertes"]:
            ws.cell(row=r, column=1, value={"danger": "Bloquant", "warn": "Vigilance", "info": "Info"}[a["niveau"]])
            ws.cell(row=r, column=2, value=a["message"])
            r += 1
    _largeurs(ws, 44, 18, 6)


ONGLETS: dict[str, tuple[str, Callable]] = {
    "hypotheses": ("1. Hypothèses", feuille_hypotheses),
    "programme": ("2. Programme et loyers", feuille_programme),
    "prix-revient": ("3. Prix de revient", feuille_prix_revient),
    "financement": ("4. Plan de financement", feuille_financement),
    "reservations": ("5. Réservations", feuille_reservations),
    "bilan-operation": ("6. Bilan d'opération", feuille_bilan_operation),
    "exploitation": ("7. Bilan d'exploitation", feuille_exploitation),
    "synthese": ("8. Synthèse", feuille_synthese),
}


PERIMETRES = {"exploitation": "7. Exploitation", "synthese": "8. Synthèse"}  # onglets déclinés avec et sans BRS (titres courts : 31 car. max)


def build_simulateur_xlsx(data: dict, res: dict, onglet: str, brs: str = "avec") -> BytesIO:
    """brs : « avec » ou « sans » pour l'export d'un onglet 7 ou 8 ; le classeur complet contient les deux."""
    wb = Workbook()
    wb.remove(wb.active)
    a_des_brs = res["synthese"]["a_des_brs"]
    feuilles: list[tuple[str, Callable, dict]] = []
    for cle in (list(ONGLETS) if onglet == "complet" else [onglet]):
        titre, fn = ONGLETS[cle]
        if cle not in PERIMETRES or not a_des_brs:
            feuilles.append((titre, fn, {"avec_brs": False} if cle == "exploitation" else {}))
        elif onglet == "complet":
            feuilles += [(f"{PERIMETRES[cle]} avec BRS", fn, {"avec_brs": True}), (f"{PERIMETRES[cle]} sans BRS", fn, {"avec_brs": False})]
        else:
            avec = brs != "sans"
            feuilles.append((f"{PERIMETRES[cle]} {'avec' if avec else 'sans'} BRS", fn, {"avec_brs": avec}))
    for titre, fn, opts in feuilles:
        ws = wb.create_sheet(titre[:31])
        fn(ws, data, res, **opts)
        ws.sheet_view.showGridLines = False
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
