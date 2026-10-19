"""Facturation : brouillons déclenchés par l'avancement des projets, émission numérotée, paiements, avoirs, journal des ventes."""
import io
import re
from datetime import date, datetime, timedelta, timezone

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel
from sqlmodel import Session, select

from ..auth import require_admin, require_saisie
from ..config import settings
from ..db import get_session
from ..models import AuditLog, Facture, Projet, User
from ..notifier import notifier
from .. import projets_outils as outils
from .projets import _admins, identite

router = APIRouter()

FACTURES_DIR = settings.data_dir / "factures"
FACTURES_DIR.mkdir(parents=True, exist_ok=True)


def _facture(session: Session, fid: int) -> Facture:
    f = session.get(Facture, fid)
    if not f:
        raise HTTPException(404, "Facture introuvable")
    return f


def en_retard(f: Facture, aujourd_hui: date | None = None) -> bool:
    return f.type == "facture" and f.statut == "emise" and f.date_echeance is not None and f.date_echeance < (aujourd_hui or date.today())


def vue(session: Session, f: Facture) -> dict:
    p = session.get(Projet, f.projet_id) if f.projet_id else None
    origine = session.get(Facture, f.facture_origine_id) if f.facture_origine_id else None
    avoir = session.exec(select(Facture).where(Facture.facture_origine_id == f.id)).first() if f.type == "facture" else None
    return {**f.model_dump(), "en_retard": en_retard(f), "projet_code": p.code if p else None, "projet_nom": p.nom if p else None,
            "numero_origine": origine.numero if origine else None,
            "avoir": {"id": avoir.id, "numero": avoir.numero, "statut": avoir.statut} if avoir else None}


def _rappels_retard(session: Session) -> None:
    """Prévient une fois les administrateurs (et le responsable du projet) quand une facture passe son échéance."""
    for f in session.exec(select(Facture).where(Facture.statut == "emise").where(Facture.rappel_retard_envoye == False)).all():  # noqa: E712
        if not en_retard(f):
            continue
        p = session.get(Projet, f.projet_id) if f.projet_id else None
        for dest in dict.fromkeys([*_admins(session), *([p.responsable] if p and p.responsable else [])]):
            notifier(session, dest, f"Facture impayée à échéance : {f.numero}",
                     f"{f.client_nom} — {outils.eur(f.total_ttc)} TTC, échue le {outils.date_fr(f.date_echeance)}.",
                     type_notif="deadline", url=f"/facturation?id={f.id}", target_type="facture", target_id=f.id)
        f.rappel_retard_envoye = True
        session.add(f)
    session.commit()


@router.get("")
def lister(session: Session = Depends(get_session), user: User = Depends(require_saisie)) -> dict:
    _rappels_retard(session)
    rows = list(session.exec(select(Facture).order_by(Facture.id.desc())).all())
    an = date.today().year
    factures = [f for f in rows if f.type == "facture"]
    return {
        "factures": [vue(session, f) for f in rows],
        "indicateurs": {
            "brouillons": sum(1 for f in rows if f.statut == "brouillon"),
            "a_emettre_ttc": round(sum(f.total_ttc for f in factures if f.statut == "brouillon"), 2),
            "en_attente_ttc": round(sum(f.total_ttc for f in factures if f.statut == "emise"), 2),
            "en_retard_ttc": round(sum(f.total_ttc for f in factures if en_retard(f)), 2),
            "nb_en_retard": sum(1 for f in factures if en_retard(f)),
            "encaisse_annee_ttc": round(sum(f.total_ttc for f in factures if f.statut == "payee" and f.date_paiement and f.date_paiement.year == an), 2),
            "facture_annee_ht": round(sum(f.total_ht for f in factures if f.statut in ("emise", "payee") and f.date_emission and f.date_emission.year == an), 2),
        },
    }


@router.get("/export.xlsx")
def exporter(session: Session = Depends(get_session), user: User = Depends(require_saisie)):
    """Journal des ventes : factures et avoirs émis, dans l'ordre des numéros."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal des ventes"
    entetes = ["Numéro", "Type", "Date", "Échéance", "Client", "SIREN", "Objet", "Projet", "Total HT", "TVA", "Total TTC", "Statut", "Payée le"]
    ws.append(entetes)
    for c in ws[1]:
        c.font = Font(bold=True)
    emises = [f for f in session.exec(select(Facture)).all() if f.numero]
    for f in sorted(emises, key=lambda x: (x.type, x.numero)):
        p = session.get(Projet, f.projet_id) if f.projet_id else None
        ws.append([f.numero, "Avoir" if f.type == "avoir" else "Facture", f.date_emission, f.date_echeance, f.client_nom, f.client_siren,
                   f.objet, p.code if p else "", f.total_ht, f.total_tva, f.total_ttc,
                   {"emise": "Émise", "payee": "Payée", "annulee": "Annulée"}.get(f.statut, f.statut), f.date_paiement])
    for col, w in zip("ABCDEFGHIJKLM", (16, 9, 12, 12, 32, 14, 44, 14, 13, 12, 13, 10, 12)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row[8:11]:
            c.number_format = '#,##0.00 "€"'
        for i in (2, 3, 12):
            row[i].number_format = "DD/MM/YYYY"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="journal_des_ventes.xlsx"'})


@router.get("/{fid}")
def lire(fid: int, session: Session = Depends(get_session), user: User = Depends(require_saisie)) -> dict:
    return vue(session, _facture(session, fid))


class FactureBody(BaseModel):
    projet_id: int | None = None
    client_nom: str | None = None
    client_adresse: str | None = None
    client_siren: str | None = None
    objet: str | None = None
    lignes: list[dict] | None = None
    notes: str | None = None


def _lignes_valides(lignes: list[dict], avoir: bool = False) -> list[dict]:
    out = []
    for l in lignes:
        d = str(l.get("designation", "")).strip()
        if not d:
            raise HTTPException(422, "Chaque ligne doit avoir une désignation")
        try:
            q, pu, tva = float(l.get("quantite", 1)), float(l.get("prix_unitaire_ht", 0)), float(l.get("taux_tva", 0.2))
        except (TypeError, ValueError):
            raise HTTPException(422, "Quantité, prix ou TVA invalide")
        if not 0 <= tva <= 0.2:
            raise HTTPException(422, "Taux de TVA invalide (entre 0 et 20 %)")
        out.append({"designation": d[:300], "quantite": q, "prix_unitaire_ht": pu, "taux_tva": tva})
    if not out:
        raise HTTPException(422, "Ajoutez au moins une ligne")
    return out


@router.post("")
def creer(body: FactureBody, session: Session = Depends(get_session), user: User = Depends(require_admin)) -> dict:
    """Facture manuelle (hors déclenchement par une étape), en brouillon."""
    p = session.get(Projet, body.projet_id) if body.projet_id else None
    if body.projet_id and not p:
        raise HTTPException(404, "Projet introuvable")
    lignes = _lignes_valides(body.lignes or [])
    ht, tva, ttc = outils.totaux(lignes)
    f = Facture(projet_id=p.id if p else None, client_nom=body.client_nom or (p.client_nom if p else ""),
                client_adresse=body.client_adresse or (p.client_adresse if p else ""), client_siren=body.client_siren or (p.client_siren if p else ""),
                objet=body.objet or (f"{p.code} — {p.nom}" if p else ""), lignes=lignes, total_ht=ht, total_tva=tva, total_ttc=ttc,
                notes=body.notes or "", cree_par=user.username)
    session.add(f)
    session.commit()
    session.refresh(f)
    return vue(session, f)


@router.put("/{fid}")
def modifier(fid: int, body: FactureBody, session: Session = Depends(get_session), user: User = Depends(require_admin)) -> dict:
    f = _facture(session, fid)
    if f.statut != "brouillon":
        raise HTTPException(409, "Une facture émise ne peut plus être modifiée : établissez un avoir")
    data = body.model_dump(exclude_unset=True, exclude={"projet_id"})
    if "lignes" in data:
        f.lignes = _lignes_valides(data.pop("lignes") or [], f.type == "avoir")
        f.total_ht, f.total_tva, f.total_ttc = outils.totaux(f.lignes)
    for k, v in data.items():
        setattr(f, k, v or "")
    f.updated_at = datetime.now(timezone.utc)
    session.add(f)
    session.commit()
    session.refresh(f)
    return vue(session, f)


@router.delete("/{fid}")
def supprimer(fid: int, session: Session = Depends(get_session), user: User = Depends(require_admin)) -> dict:
    f = _facture(session, fid)
    if f.statut != "brouillon":
        raise HTTPException(409, "Une facture émise ne peut pas être supprimée : établissez un avoir")
    session.delete(f)
    session.add(AuditLog(action="facture_suppression_brouillon", target_type="facture", target_id=fid, user=user.username,
                         payload={"objet": f.objet, "total_ttc": f.total_ttc}))
    session.commit()
    return {"ok": True}


def _prochain_numero(session: Session, prefixe: str, an: int) -> str:
    n = 0
    for num in session.exec(select(Facture.numero).where(Facture.numero.like(f"{prefixe}-{an}-%"))).all():
        m = re.search(r"-(\d+)$", num or "")
        if m:
            n = max(n, int(m.group(1)))
    return f"{prefixe}-{an}-{n + 1:04d}"


class EmettreBody(BaseModel):
    date_emission: date | None = None


@router.post("/{fid}/emettre")
def emettre(fid: int, body: EmettreBody | None = None, session: Session = Depends(get_session), user: User = Depends(require_admin)) -> dict:
    """Attribue le numéro (séquence continue par année), fige le PDF et rend la facture non modifiable."""
    f = _facture(session, fid)
    if f.statut != "brouillon":
        raise HTTPException(409, "Cette facture est déjà émise")
    if not f.client_nom.strip():
        raise HTTPException(422, "Renseignez le client avant d'émettre la facture")
    if not f.lignes:
        raise HTTPException(422, "La facture n'a aucune ligne")
    ident = identite(session)
    jour = (body.date_emission if body and body.date_emission else None) or date.today()
    if jour > date.today():
        raise HTTPException(422, "La date d'émission ne peut pas être dans le futur")
    prefixe = ident.get("prefixe_avoir" if f.type == "avoir" else "prefixe_facture") or ("AV" if f.type == "avoir" else "FC")
    dernier = session.exec(select(Facture).where(Facture.numero.like(f"{prefixe}-{jour.year}-%")).where(Facture.date_emission > jour)).first()
    if dernier:
        raise HTTPException(422, f"La facture {dernier.numero} est datée du {outils.date_fr(dernier.date_emission)} : "
                                 "la numérotation doit suivre l'ordre chronologique")
    f.numero = _prochain_numero(session, prefixe, jour.year)
    f.date_emission = jour
    f.date_echeance = jour + timedelta(days=int(ident.get("delai_paiement_jours") or 30)) if f.type == "facture" else None
    f.statut = "emise"
    f.emise_par = user.username
    origine = session.get(Facture, f.facture_origine_id) if f.facture_origine_id else None
    pdf = outils.facture_pdf(f, ident, origine.numero if origine else None)
    chemin = FACTURES_DIR / f"{f.numero}.pdf"
    chemin.write_bytes(pdf)
    f.chemin_pdf = chemin.name
    f.updated_at = datetime.now(timezone.utc)
    session.add(f)
    if origine:  # avoir total : la facture d'origine est annulée
        origine.statut = "annulee"
        origine.updated_at = datetime.now(timezone.utc)
        session.add(origine)
    session.add(AuditLog(action="facture_emission", target_type="facture", target_id=f.id, user=user.username,
                         payload={"numero": f.numero, "total_ttc": f.total_ttc, "type": f.type}))
    session.commit()
    session.refresh(f)
    return vue(session, f)


class PaiementBody(BaseModel):
    date_paiement: date | None = None
    mode_paiement: str = "virement"


@router.post("/{fid}/payer")
def payer(fid: int, body: PaiementBody, session: Session = Depends(get_session), user: User = Depends(require_admin)) -> dict:
    f = _facture(session, fid)
    if f.type != "facture" or f.statut != "emise":
        raise HTTPException(409, "Seule une facture émise et non réglée peut être marquée payée")
    f.statut = "payee"
    f.date_paiement = body.date_paiement or date.today()
    f.mode_paiement = body.mode_paiement[:40]
    f.updated_at = datetime.now(timezone.utc)
    session.add(f)
    p = session.get(Projet, f.projet_id) if f.projet_id else None
    if p and p.responsable:
        notifier(session, p.responsable, f"Facture réglée : {f.numero}", f"{f.client_nom} — {outils.eur(f.total_ttc)} TTC.",
                 type_notif="info", url=f"/projets/{p.id}", emise_par=user.username, target_type="facture", target_id=f.id)
    session.add(AuditLog(action="facture_paiement", target_type="facture", target_id=f.id, user=user.username,
                         payload={"numero": f.numero, "date": f.date_paiement.isoformat()}))
    session.commit()
    session.refresh(f)
    return vue(session, f)


@router.post("/{fid}/avoir")
def avoir(fid: int, session: Session = Depends(get_session), user: User = Depends(require_admin)) -> dict:
    """Prépare un avoir (brouillon) qui annule la facture en totalité."""
    f = _facture(session, fid)
    if f.type != "facture" or f.statut not in ("emise", "payee"):
        raise HTTPException(409, "Seule une facture émise peut être annulée par un avoir")
    if session.exec(select(Facture).where(Facture.facture_origine_id == f.id)).first():
        raise HTTPException(409, "Un avoir existe déjà pour cette facture")
    lignes = [{**l, "prix_unitaire_ht": -float(l["prix_unitaire_ht"])} for l in f.lignes]
    ht, tva, ttc = outils.totaux(lignes)
    a = Facture(type="avoir", facture_origine_id=f.id, projet_id=f.projet_id, etape_id="", client_nom=f.client_nom,
                client_adresse=f.client_adresse, client_siren=f.client_siren, objet=f"Avoir sur la facture {f.numero} — {f.objet}",
                lignes=lignes, total_ht=ht, total_tva=tva, total_ttc=ttc, cree_par=user.username)
    session.add(a)
    session.commit()
    session.refresh(a)
    return vue(session, a)


@router.get("/{fid}/pdf")
def pdf(fid: int, session: Session = Depends(get_session), user: User = Depends(require_saisie)) -> Response:
    f = _facture(session, fid)
    if f.chemin_pdf and (FACTURES_DIR / f.chemin_pdf).exists():
        data = (FACTURES_DIR / f.chemin_pdf).read_bytes()  # PDF figé à l'émission
    else:
        origine = session.get(Facture, f.facture_origine_id) if f.facture_origine_id else None
        data = outils.facture_pdf(f, identite(session), origine.numero if origine else None)
    nom = f"{f.numero or f'brouillon_{f.id}'}.pdf"
    return Response(data, media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{nom}"'})
